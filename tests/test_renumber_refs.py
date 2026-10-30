#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""renumber_refs.py 的回归测试。

运行：
    python -m unittest discover -s tests -v
"""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

import docx
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import renumber_refs as rr  # noqa: E402


def cite_para(document, text, cites):
    """写一段正文：cites 里给出需要排上标的片段（按出现顺序）。"""
    p = document.add_paragraph()
    rest = text
    for c in cites:
        head, rest = rest.split(c, 1)
        if head:
            p.add_run(head)
        r = p.add_run(c)
        r.font.superscript = True
    if rest:
        p.add_run(rest)
    return p


def make_doc(path: Path) -> docx.Document:
    d = docx.Document()
    cite_para(d, "本文首先研究了A问题[3]。", ["[3]"])
    cite_para(d, "随后讨论了B问题[1,2]。", ["[1,2]"])
    p = d.add_paragraph()
    p.add_run("晶向[001]应保持在正文中，不上标。")       # 非上标，必须被忽略
    cite_para(d, "再次引用C问题[1]与D问题[5]。", ["[1]", "[5]"])
    cite_para(d, "最后是E问题[4]。", ["[4]"])
    d.add_paragraph("参考文献")
    for i, name in enumerate("ABCDE", start=1):
        d.add_paragraph("[%d] %s et al., Journal, 1 (2020) 1-2." % (i, name))
    d.save(path)
    return d


class RenumberRefsTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "thesis.docx"
        make_doc(self.path)

    def tearDown(self):
        self.tmp.cleanup()

    def test_mapping_follows_first_citation(self):
        mapping = rr.renumber(str(self.path), dry_run=True)
        # 首次出现顺序：3, 1, 2, 5, 4
        self.assertEqual(mapping, {3: 1, 1: 2, 2: 3, 5: 4, 4: 5})

    def test_tokens_and_list_are_rewritten(self):
        rr.renumber(str(self.path))
        d = docx.Document(str(self.path))
        texts = [p.text for p in d.paragraphs]
        self.assertEqual(texts[0], "本文首先研究了A问题[1]。")
        self.assertEqual(texts[1], "随后讨论了B问题[2-3]。")
        self.assertEqual(texts[2], "晶向[001]应保持在正文中，不上标。")
        self.assertEqual(texts[3], "再次引用C问题[2]与D问题[4]。")
        self.assertEqual(texts[4], "最后是E问题[5]。")
        idx = texts.index("参考文献")
        entries = texts[idx + 1: idx + 6]
        self.assertEqual([e.split()[0] for e in entries], ["[1]", "[2]", "[3]", "[4]", "[5]"])
        # 新表顺序 = 旧 3,1,2,5,4 → 条目内容依次为 C, A, B, E, D
        self.assertEqual([e.split()[1] for e in entries], ["C", "A", "B", "E", "D"])

    def test_non_superscript_brackets_are_ignored(self):
        mapping = rr.renumber(str(self.path), dry_run=True)
        self.assertNotIn(1, [k for k in mapping if k not in (1, 2, 3, 4, 5)])  # 只为可读性
        self.assertEqual(sorted(mapping), [1, 2, 3, 4, 5])

    def test_multi_run_token_is_rewritten(self):
        d = docx.Document(str(self.path))
        p = d.add_paragraph()
        r1 = p.add_run("[2,")
        r1.font.superscript = True
        r2 = p.add_run("4]")
        r2.font.superscript = True
        d.save(str(self.path))
        rr.renumber(str(self.path))
        d2 = docx.Document(str(self.path))
        self.assertTrue(any(p.text.endswith("[3,5]。") or "[3,5]" in p.text for p in d2.paragraphs))

    def test_uncited_reference_is_reported(self):
        d = docx.Document(str(self.path))
        d.add_paragraph("[6] F et al., Journal, 2 (2021) 3-4.")
        d.save(str(self.path))
        with self.assertRaises(SystemExit):
            rr.renumber(str(self.path), dry_run=True)


if __name__ == "__main__":
    unittest.main()
