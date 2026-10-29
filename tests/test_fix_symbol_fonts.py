#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""fix_symbol_fonts.py 的回归测试。

运行：
    python -m unittest discover -s tests -v
"""

from __future__ import annotations

import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import docx
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
SCRIPT = ROOT / "scripts" / "fix_symbol_fonts.py"
sys.path.insert(0, str(ROOT / "scripts"))
import fix_symbol_fonts as fsf  # noqa: E402


def make_doc(path: Path, text: str, eastasia: str = "黑体", hint: bool = True) -> docx.Document:
    """造一个带 w:hint="eastAsia" 的题注式段落。"""
    document = docx.Document()
    p = document.add_paragraph()
    run = p.add_run(text)
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), eastasia)
    if hint:
        rfonts.set(qn("w:hint"), "eastAsia")
    document.save(str(path))
    return document


def runs_text(p) -> list[str]:
    return [r.text for r in p.runs]


class UnitTest(unittest.TestCase):
    def test_is_symbol_classification(self):
        for ch in "δγ℃°±×÷′″₂μΔΣ≈≤→":
            self.assertTrue(fsf.is_symbol(ch), ch)
        for ch in "图合金（）、，。A1":
            self.assertFalse(fsf.is_symbol(ch), ch)

    def test_scan_finds_hinted_symbol_run(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "a.docx"
            document = make_doc(path, "图4.4 氧化物/金属界面处δ相的TEM分析")
            runs, samples = fsf.scan(document)
            self.assertEqual(len(runs), 1)
            self.assertTrue(samples and "δ" in samples[0])

    def test_split_keeps_text_and_cjk_font(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "b.docx"
            document = make_doc(path, "图4.4 氧化物/金属界面处δ相的TEM分析")
            p = document.paragraphs[0]
            before = p.text
            runs, _ = fsf.scan(document)
            fsf.fix_runs(runs)
            self.assertEqual(p.text, before)          # 文字一字不改
            # δ 单独成段并去掉 hint
            sym = [r for r in p.runs if "δ" in r.text]
            self.assertTrue(sym)
            for r in sym:
                rf = r._element.rPr.rFonts
                self.assertIsNone(rf.get(qn("w:hint")))
                self.assertEqual(rf.get(qn("w:eastAsia")), "黑体")  # 中文字体保留
            # 中文段仍保留 hint=eastAsia，避免全角标点/引号被换成西文字体
            cn = [r for r in p.runs if "氧化物" in r.text]
            self.assertTrue(cn)
            self.assertEqual(cn[0]._element.rPr.rFonts.get(qn("w:hint")), "eastAsia")

    def test_all_symbol_run_only_drops_hint(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "c.docx"
            document = make_doc(path, "δ℃γ″")
            p = document.paragraphs[0]
            runs, _ = fsf.scan(document)
            changed, _ = fsf.fix_runs(runs)
            self.assertEqual(changed, 1)
            self.assertEqual(len(p.runs), 1)           # 全是符号，不需要切
            self.assertIsNone(p.runs[0]._element.rPr.rFonts.get(qn("w:hint")))

    def test_ascii_slash_stays_in_cjk_run(self):
        """斜杠是 ASCII，本来就按 ascii 字体排，不必跟着符号切。"""
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "c2.docx"
            document = make_doc(path, "γ′/γ″")
            p = document.paragraphs[0]
            runs, _ = fsf.scan(document)
            fsf.fix_runs(runs)
            self.assertEqual(runs_text(p), ["γ′", "/", "γ″"])
            self.assertIsNone(p.runs[0]._element.rPr.rFonts.get(qn("w:hint")))
            self.assertEqual(p.runs[1]._element.rPr.rFonts.get(qn("w:hint")), "eastAsia")

    def test_run_without_hint_untouched(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "d.docx"
            document = make_doc(path, "850 ℃氧化", hint=False)
            runs, _ = fsf.scan(document)
            self.assertEqual(runs, [])

    def test_space_preserved_after_split(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "e.docx"
            document = make_doc(path, "图4.5 无γ′/γ″层的TEM分析")
            p = document.paragraphs[0]
            before = p.text
            runs, _ = fsf.scan(document)
            fsf.fix_runs(runs)
            self.assertEqual(p.text, before)
            self.assertTrue(p.text.startswith("图4.5 无"))

    def test_check_mode_does_not_write(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "f.docx"
            make_doc(path, "δ相")
            before = path.stat().st_mtime_ns
            fsf.main([str(path), "--quiet"])
            self.assertEqual(path.stat().st_mtime_ns, before)


class CliTest(unittest.TestCase):
    def test_cli_out_writes_fixed_copy(self):
        with tempfile.TemporaryDirectory() as td:
            src = Path(td) / "src.docx"
            out = Path(td) / "fixed.docx"
            make_doc(src, "表4.2 γ″相和δ相中Nb位点的O₂吸附能")
            result = subprocess.run(
                [sys.executable, "-X", "utf8", str(SCRIPT), str(src), "--out", str(out)],
                capture_output=True, text=True, encoding="utf-8", errors="replace",
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            fixed = docx.Document(str(out))
            p = fixed.paragraphs[0]
            sym = [r for r in p.runs if "δ" in r.text]
            self.assertTrue(sym)
            self.assertIsNone(sym[0]._element.rPr.rFonts.get(qn("w:hint")))
            # 源文件不动
            raw = docx.Document(str(src))
            self.assertEqual(raw.paragraphs[0].runs[0]._element.rPr.rFonts.get(qn("w:hint")), "eastAsia")

    def test_cli_rejects_both_out_and_in_place(self):
        with tempfile.TemporaryDirectory() as td:
            src = Path(td) / "src.docx"
            make_doc(src, "δ相")
            result = subprocess.run(
                [sys.executable, "-X", "utf8", str(SCRIPT), str(src),
                 "--out", str(Path(td) / "x.docx"), "--in-place"],
                capture_output=True, text=True, encoding="utf-8", errors="replace",
            )
            self.assertNotEqual(result.returncode, 0)


if __name__ == "__main__":
    unittest.main(verbosity=2)
