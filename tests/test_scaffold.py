#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""thesis_scaffold.py 的回归测试。

运行：
    python -m unittest discover -s tests -v
"""

from __future__ import annotations

import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import docx
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
SCAFFOLD = ROOT / "scripts" / "thesis_scaffold.py"
LINT = ROOT / "scripts" / "thesis_lint.py"


def build(path: Path, *extra: str) -> subprocess.CompletedProcess:
    cmd = [
        sys.executable, "-X", "utf8", str(SCAFFOLD),
        "--out", str(path), "--title", "面向测试的示例研究",
        *extra,
    ]
    return subprocess.run(cmd, capture_output=True, text=True,
                          encoding="utf-8", errors="replace")


class ScaffoldTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls._tmp = tempfile.TemporaryDirectory()
        cls.tmp = Path(cls._tmp.name)
        cls.path = cls.tmp / "skeleton.docx"
        result = build(cls.path)
        assert result.returncode == 0, result.stdout + result.stderr
        cls.doc = docx.Document(str(cls.path))

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()

    def test_page_size_and_margins(self):
        for section in self.doc.sections:
            self.assertAlmostEqual(section.page_width.cm, 21.0, places=1)
            self.assertAlmostEqual(section.page_height.cm, 29.7, places=1)
            self.assertAlmostEqual(section.top_margin.cm, 2.2, places=1)
            self.assertAlmostEqual(section.bottom_margin.cm, 2.2, places=1)
            self.assertAlmostEqual(section.left_margin.cm, 2.5, places=1)
            self.assertAlmostEqual(section.right_margin.cm, 2.5, places=1)
            self.assertAlmostEqual(section.header_distance.cm, 1.1, places=1)
            self.assertAlmostEqual(section.footer_distance.cm, 1.1, places=1)

    def test_six_sections_with_roman_then_arabic_numbering(self):
        sections = self.doc.sections
        self.assertEqual(len(sections), 6)

        def pg(section):
            el = section._sectPr.find(qn("w:pgNumType"))
            if el is None:
                return None
            return el.get(qn("w:fmt")), el.get(qn("w:start"))

        self.assertIsNone(pg(sections[0]), "封面节不应有页码设置")
        self.assertEqual(pg(sections[1]), ("upperRoman", "1"))
        # 摘要 / 英文摘要 / 目录：必须续用罗马数字，不能各自从 1 开始
        for i in (2, 3, 4):
            self.assertIsNone(pg(sections[i]), f"第 {i} 节不应重设页码")
        self.assertEqual(pg(sections[5]), ("decimal", "1"))

    def test_odd_even_headers_and_page_field(self):
        body = self.doc.sections[-1]
        odd = "\n".join(p.text for p in body.header.paragraphs).strip()
        even = "\n".join(p.text for p in body.even_page_header.paragraphs).strip()
        self.assertEqual(odd, "博士学位论文")
        self.assertIn("面向测试的示例研究", even)

        footer_xml = body.footer._element.xml
        self.assertIn("PAGE", footer_xml)

        # 页眉下边框：双线样式之一
        found = False
        for p in list(body.header.paragraphs) + list(body.even_page_header.paragraphs):
            ppr = p._p.pPr
            if ppr is None:
                continue
            pbdr = ppr.find(qn("w:pBdr"))
            if pbdr is None:
                continue
            bottom = pbdr.find(qn("w:bottom"))
            if bottom is not None and "Thick" in (bottom.get(qn("w:val")) or ""):
                found = True
        self.assertTrue(found, "页眉应有双线边框（thinThickLargeGap 等）")

    def test_front_matter_content(self):
        text = "\n".join(p.text for p in self.doc.paragraphs)
        for needle in (
            "学校代号", "博士学位论文", "学位申请人姓名", "导师姓名及职称",
            "湖南大学博士学位论文", "学位论文原创性声明", "学位论文版权使用授权书",
            "摘   要", "关键词：", "Abstract", "Key Words:",
            "目   录", "插图索引", "附表索引",
            "结   论", "参考文献", "附录A", "致   谢",
        ):
            self.assertIn(needle, text, f"缺少：{needle}")

    def test_required_fields(self):
        instr = " ".join(
            el.text or ""
            for el in self.doc.element.body.iter(qn("w:instrText"))
        )
        self.assertIn('TOC \\o "1-4"', instr)
        self.assertIn('\\c "图"', instr)
        self.assertIn('\\c "表"', instr)
        self.assertIn("STYLEREF 1", instr)
        self.assertIn("SEQ 图", instr)
        self.assertIn("SEQ 表", instr)
        self.assertIn("REF _RefFigExample", instr)
        self.assertIn("REF _RefTabExample", instr)

    def test_heading_numbering_is_automatic(self):
        found = 0
        for p in self.doc.paragraphs:
            if p.style.name == "Heading 1":
                ppr = p._p.pPr
                if ppr is not None and ppr.find(qn("w:numPr")) is not None:
                    num = ppr.find(qn("w:numPr")).find(qn("w:numId"))
                    self.assertIn(num.get(qn("w:val")), {"0", "90"})
                    if num.get(qn("w:val")) == "90":
                        found += 1
        self.assertGreaterEqual(found, 1, "第 1 章应使用自动编号")

    def test_lint_reports_no_error(self):
        result = subprocess.run(
            [sys.executable, "-X", "utf8", str(LINT), str(self.path),
             "--degree", "doctor"],
            capture_output=True, text=True, encoding="utf-8", errors="replace",
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertIn("ERROR 0", result.stdout)

    def test_master_degree_labels(self):
        path = self.tmp / "master.docx"
        result = build(path, "--degree", "master")
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        document = docx.Document(str(path))
        text = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("硕士学位论文", text)
        self.assertIn("湖南大学硕士学位论文", text)


if __name__ == "__main__":
    unittest.main(verbosity=2)
