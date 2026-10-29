#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""fix_latin_fonts.py 的回归测试。

运行：
    python -m unittest discover -s tests -v
"""

from __future__ import annotations

import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import docx
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
SCRIPT = ROOT / "scripts" / "fix_latin_fonts.py"
sys.path.insert(0, str(ROOT / "scripts"))
import fix_latin_fonts as flf  # noqa: E402


def make_doc(path: Path, text: str = "IN718 合金在 850 ℃ 氧化") -> docx.Document:
    document = docx.Document()
    p = document.add_paragraph()
    run = p.add_run(text)
    run.font.name = "宋体"
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "宋体")
    document.save(str(path))
    return document


class UnitTest(unittest.TestCase):
    def test_scan_finds_latin_in_cjk_font(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "a.docx"
            document = make_doc(path)
            runs, samples = flf.scan(document)
            self.assertEqual(len(runs), 1)
            self.assertTrue(samples and "IN718" in samples[0])

    def test_fix_changes_ascii_only(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "b.docx"
            document = make_doc(path)
            runs, _ = flf.scan(document)
            self.assertEqual(flf.fix_runs(runs), 1)
            rfonts = runs[0].find(qn("w:rPr")).find(qn("w:rFonts"))
            self.assertEqual(rfonts.get(qn("w:ascii")), "Times New Roman")
            self.assertEqual(rfonts.get(qn("w:hAnsi")), "Times New Roman")
            # 中文字体必须保持不变
            self.assertEqual(rfonts.get(qn("w:eastAsia")), "宋体")

    def test_fix_style_keeps_eastasia(self):
        document = docx.Document()
        normal = document.styles["Normal"]
        normal.font.name = "宋体"
        normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
        document.add_paragraph("1.1 研究背景")
        changed = flf.fix_styles(document)
        self.assertIn("Normal", changed)
        rfonts = normal.element.rPr.rFonts
        self.assertEqual(rfonts.get(qn("w:ascii")), "Times New Roman")
        self.assertEqual(rfonts.get(qn("w:eastAsia")), "宋体")

    def test_already_latin_font_is_untouched(self):
        document = docx.Document()
        run = document.add_paragraph().add_run("IN718")
        run.font.name = "Times New Roman"
        runs, _ = flf.scan(document)
        self.assertEqual(runs, [])

    def test_check_mode_does_not_write(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "c.docx"
            make_doc(path)
            before = path.stat().st_mtime_ns
            flf.main([str(path), "--quiet"])
            self.assertEqual(path.stat().st_mtime_ns, before)


class CliTest(unittest.TestCase):
    def test_cli_out_writes_fixed_copy(self):
        with tempfile.TemporaryDirectory() as td:
            src = Path(td) / "src.docx"
            out = Path(td) / "fixed.docx"
            make_doc(src)
            result = subprocess.run(
                [sys.executable, "-X", "utf8", str(SCRIPT), str(src), "--out", str(out)],
                capture_output=True, text=True, encoding="utf-8", errors="replace",
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertTrue(out.exists())
            fixed = docx.Document(str(out))
            rfonts = fixed.paragraphs[0].runs[0]._element.rPr.rFonts
            self.assertEqual(rfonts.get(qn("w:ascii")), "Times New Roman")
            self.assertEqual(rfonts.get(qn("w:eastAsia")), "宋体")
            # 源文件不应被改动
            raw = docx.Document(str(src))
            self.assertEqual(raw.paragraphs[0].runs[0]._element.rPr.rFonts.get(qn("w:ascii")), "宋体")

    def test_cli_rejects_both_out_and_in_place(self):
        with tempfile.TemporaryDirectory() as td:
            src = Path(td) / "src.docx"
            make_doc(src)
            result = subprocess.run(
                [sys.executable, "-X", "utf8", str(SCRIPT), str(src),
                 "--out", str(Path(td) / "x.docx"), "--in-place"],
                capture_output=True, text=True, encoding="utf-8", errors="replace",
            )
            self.assertNotEqual(result.returncode, 0)


if __name__ == "__main__":
    unittest.main(verbosity=2)
