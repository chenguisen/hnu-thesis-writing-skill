#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""thesis_lint.py 的回归测试。

运行：
    python -m unittest discover -s tests -v
或：
    python tests/test_lint.py
"""

from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import docx
from docx.oxml import OxmlElement, parse_xml
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent
LINT = ROOT / "scripts" / "thesis_lint.py"
FIXTURE = Path(__file__).resolve().parent / "fixture_master.docx"
sys.path.insert(0, str(ROOT / "scripts"))
import thesis_lint as tl  # noqa: E402


def run_lint(path: Path, *extra: str) -> subprocess.CompletedProcess:
    cmd = [sys.executable, "-X", "utf8", str(LINT), str(path), "--degree", "master", *extra]
    return subprocess.run(
        cmd, capture_output=True, text=True, encoding="utf-8", errors="replace"
    )


class UnitTest(unittest.TestCase):
    def test_heading_without_space_after_number(self):
        self.assertEqual(tl.classify_heading("1.1研究背景"), "l2")
        self.assertEqual(tl.classify_heading("1.1.1研究现状"), "l3")
        self.assertEqual(tl.classify_heading("1.1.1.1国外现状"), "l4")

    def test_toc_tab_leader_is_not_heading(self):
        self.assertTrue(tl.is_toc_line("第1章 绪论\t1"))
        self.assertTrue(tl.is_toc_line("1.1 研究背景\t2"))
        self.assertIsNone(tl.classify_heading("1.1 研究背景\t2"))

    def test_toc_style_is_not_heading(self):
        self.assertTrue(tl.is_toc_line("1.1 研究背景", "TOC 2"))
        self.assertIsNone(tl.classify_heading("第1章 绪论", "TOC 1"))

    def test_explicit_nil_table_borders_override_grid_style(self):
        document = docx.Document()
        table = document.add_table(rows=1, cols=1)
        table.style = "Table Grid"
        self.assertTrue(tl.table_has_side_borders(table))

        # 显式把左右边线设为 nil 后，不应再按 Table Grid 误报
        tbl_pr = table._tbl.tblPr
        borders = tbl_pr.find(tl.qn("w:tblBorders"))
        if borders is None:
            from docx.oxml import OxmlElement

            borders = OxmlElement("w:tblBorders")
            tbl_pr.append(borders)
        for side in ("left", "right"):
            el = borders.find(tl.qn("w:" + side))
            if el is None:
                from docx.oxml import OxmlElement

                el = OxmlElement("w:" + side)
                borders.append(el)
            el.set(tl.qn("w:val"), "nil")
        self.assertFalse(tl.table_has_side_borders(table))

    def test_paragraph_format_from_style(self):
        document = docx.Document()
        normal = document.styles["Normal"]
        normal.paragraph_format.first_line_indent = docx.shared.Pt(24)
        normal.paragraph_format.line_spacing = docx.shared.Pt(20)
        p = document.add_paragraph("测试段落" * 10)
        self.assertEqual(tl.effective_first_line_indent(p), docx.shared.Pt(24))
        self.assertEqual(tl.effective_line_spacing(p), docx.shared.Pt(20))

    def test_count_ref_fields(self):
        document = docx.Document()
        p = document.add_paragraph()
        run = p.add_run()
        instr = OxmlElement("w:instrText")
        instr.text = r" REF _Ref12345 \h "
        run._r.append(instr)
        self.assertEqual(tl.count_ref_fields(document), 1)

    def test_auto_numbered_heading_recognised(self):
        # Word 自动编号时正文里只有“绪论”，没有“第1章”
        self.assertEqual(tl.style_heading_level("Heading 1"), 1)
        self.assertEqual(tl.style_heading_level("标题 2"), 2)
        self.assertIsNone(tl.style_heading_level("Normal"))
        self.assertEqual(
            tl.classify_heading("绪论", "Heading 1", auto_level=1), "chapter"
        )
        self.assertEqual(
            tl.classify_heading("研究背景", "Heading 2", auto_level=2), "l2"
        )

    def test_thick_thin_header_border_counts_as_double(self):
        document = docx.Document()
        section = document.sections[0]
        p = section.header.paragraphs[0]
        ppr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(tl.qn("w:val"), "thinThickLargeGap")
        pbdr.append(bottom)
        ppr.append(pbdr)
        self.assertTrue(tl.header_has_double_border(section))

        bottom.set(tl.qn("w:val"), "single")
        self.assertFalse(tl.header_has_double_border(section))

    def test_pgnumtype_without_fmt_is_decimal(self):
        document = docx.Document()
        section = document.sections[0]
        el = OxmlElement("w:pgNumType")
        el.set(tl.qn("w:start"), "1")
        section._sectPr.append(el)
        self.assertEqual(tl.get_section_start_type(section), "decimal")

    def test_latin_in_cjk_font_counts_as_usage(self):
        document = docx.Document()
        run = document.add_paragraph().add_run("IN718 合金 850 h")
        run.font.name = "宋体"
        run._element.get_or_add_rPr().get_or_add_rFonts().set(
            tl.qn("w:eastAsia"), "宋体"
        )
        runs, samples, _styles = tl.latin_in_cjk_font_usage(document)
        self.assertEqual(runs, 1)
        self.assertTrue(samples and "IN718" in samples[0])

    def test_orphan_figure_is_reported(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "orphan.docx"
            document = docx.Document(str(FIXTURE))
            document.add_paragraph("图1.1  未被正文引用的示意图")
            document.save(str(path))
            result = run_lint(path)
            self.assertEqual(result.returncode, 0, result.stdout)
            self.assertIn("XREF-02", result.stdout)


class LintTest(unittest.TestCase):
    def run_lint(self, path: Path, *extra: str) -> subprocess.CompletedProcess:
        return run_lint(path, *extra)

    def test_compliant_fixture_has_no_error(self):
        result = self.run_lint(FIXTURE)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertIn("ERROR 0", result.stdout)

    def test_non_a4_page_is_an_error(self):
        with tempfile.TemporaryDirectory() as td:
            bad = Path(td) / "bad.docx"
            document = docx.Document(str(FIXTURE))
            section = document.sections[0]
            section.page_width = docx.shared.Cm(21.59)  # Letter
            section.page_height = docx.shared.Cm(27.94)
            document.save(str(bad))
            result = self.run_lint(bad)
            self.assertEqual(result.returncode, 1, result.stdout)
            self.assertIn("PAGE-02", result.stdout)

    def test_toc_tab_leader_does_not_trigger_heading_warnings(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "with_toc.docx"
            document = docx.Document(str(FIXTURE))
            for text in (
                "第1章 绪论\t1",
                "1.1 研究背景与意义\t2",
                "1.1.1 国内外研究现状\t3",
                "第2章 理论基础\t4",
                "参考文献\t5",
            ):
                document.add_paragraph(text)
            document.save(str(path))
            result = self.run_lint(path)
            self.assertEqual(result.returncode, 0, result.stdout)
            self.assertNotIn("HEAD-04", result.stdout)
            self.assertNotIn("HEAD-10", result.stdout)
            self.assertNotIn("CAP-02", result.stdout)

    def test_chapter_mode_skips_front_matter_checks(self):
        result = self.run_lint(FIXTURE, "--chapter")
        self.assertEqual(result.returncode, 0, result.stdout)
        self.assertNotIn("[FRONT-01]", result.stdout)
        self.assertNotIn("[ABS-01]", result.stdout)
        self.assertNotIn("[REF-01]", result.stdout)

    def test_loose_figure_and_placeholder_table_are_flagged(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "loose.docx"
            document = docx.Document(str(FIXTURE))
            document.add_paragraph("图1  测试示意图")
            document.add_paragraph("如图1所示，测试正文。")
            document.add_paragraph("表x. 测试表")
            document.save(str(path))
            result = self.run_lint(path, "--chapter")
            self.assertEqual(result.returncode, 0, result.stdout)
            self.assertIn("CAP-03", result.stdout)
            self.assertIn("CAP-04", result.stdout)

    def test_duplicated_chapter_number_in_caption_is_warned(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "dupchap.docx"
            document = docx.Document(str(FIXTURE))
            document.add_paragraph("图4.4.2  氧化 1 h 后样品的表面与截面形貌及元素分布")
            document.add_paragraph("表2.2.1  测试表格")
            document.save(str(path))
            result = self.run_lint(path, "--chapter")
            self.assertEqual(result.returncode, 0, result.stdout)
            self.assertIn("CAP-08", result.stdout)
            self.assertIn("重复章号", result.stdout)

    def test_latin_in_cjk_font_is_warned(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "latin.docx"
            document = docx.Document(str(FIXTURE))
            for i in range(6):
                run = document.add_paragraph().add_run(f"IN718 850 h 试样 {i}")
                run.font.name = "宋体"
                run._element.get_or_add_rPr().get_or_add_rFonts().set(
                    tl.qn("w:eastAsia"), "宋体"
                )
            document.save(str(path))
            result = run_lint(path, "--chapter")
            self.assertEqual(result.returncode, 0, result.stdout)
            self.assertIn("FONT-07", result.stdout)
            self.assertIn("Times New Roman", result.stdout)

    def test_body_heiti_is_warned(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "heiti.docx"
            document = docx.Document(str(FIXTURE))
            for _ in range(3):
                p = document.add_paragraph()
                run = p.add_run("图1.1")
                run._element.get_or_add_rPr().get_or_add_rFonts().set(
                    tl.qn("w:eastAsia"), "黑体"
                )
            document.save(str(path))
            result = run_lint(path, "--chapter")
            self.assertEqual(result.returncode, 0, result.stdout)
            self.assertIn("FONT-08", result.stdout)

    def test_symbol_in_cjk_font_is_warned(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "symbol.docx"
            document = docx.Document(str(FIXTURE))
            for text in ("图4.4 氧化物/金属界面处δ相的TEM分析",
                         "图4.5 无γ′/γ″层的TEM分析",
                         "表4.2 γ″相和δ相中Nb位点的O₂吸附能"):
                run = document.add_paragraph().add_run(text)
                rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
                rfonts.set(tl.qn("w:eastAsia"), "黑体")
                rfonts.set(tl.qn("w:hint"), "eastAsia")
            document.save(str(path))
            result = run_lint(path, "--chapter")
            self.assertEqual(result.returncode, 0, result.stdout)
            self.assertIn("FONT-09", result.stdout)
            self.assertIn("fix_symbol_fonts.py", result.stdout)

    def _caption_doc(self, path: Path) -> docx.Document:
        """在样例文档末尾加一条 Caption 样式的表题（样式设成五号）。"""
        document = docx.Document(str(FIXTURE))
        style = document.styles["Caption"]
        style.font.size = Pt(10.5)
        document.add_paragraph("表9.9  测试题注", style="Caption")
        document.save(str(path))
        return document

    def test_caption_size_override_is_warned(self):
        """题注样式是五号，但 run 上写死了 12 pt：渲染会变成小四。"""
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "capsize.docx"
            document = self._caption_doc(path)
            cap = document.paragraphs[-1]
            for r in cap.runs:
                r.font.size = Pt(12)
            document.save(str(path))
            result = run_lint(path, "--chapter")
            self.assertEqual(result.returncode, 0, result.stdout)
            self.assertIn("FONT-10", result.stdout)

    def test_caption_keep_flags_are_warned(self):
        """题注没设“段中不分页”，会被拆到两页。"""
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "capkeep.docx"
            document = self._caption_doc(path)
            cap = document.paragraphs[-1]
            ppr = cap._element.get_or_add_pPr()
            for tag in ("keepLines", "keepNext"):
                for el in ppr.findall(tl.qn("w:" + tag)):
                    ppr.remove(el)
            document.save(str(path))
            result = run_lint(path, "--chapter")
            self.assertEqual(result.returncode, 0, result.stdout)
            self.assertIn("CAP-06", result.stdout)

    def test_caption_inline_math_is_warned(self):
        """题注里夹行内公式对象：会排到题注文字前，并被索引域复制。"""
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "capmath.docx"
            document = self._caption_doc(path)
            cap = document.paragraphs[-1]
            omml = (
                '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
                '<m:r><m:t>(10-10)</m:t></m:r></m:oMath>'
            )
            cap._element.insert(1, parse_xml(omml))
            document.save(str(path))
            result = run_lint(path, "--chapter")
            self.assertEqual(result.returncode, 0, result.stdout)
            self.assertIn("CAP-05", result.stdout)

    def test_figure_note_inline_math_is_warned(self):
        """图注（题注下一段）里夹行内公式：晶面指数会变成 CambriaMath 对象。"""
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "notemath.docx"
            document = self._caption_doc(path)
            note = document.add_paragraph("(a) (0001)面Cr终止；(b) 面O终止。")
            omml = (
                '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
                '<m:r><m:t>(10-10)</m:t></m:r></m:oMath>'
            )
            note._element.insert(1, parse_xml(omml))
            document.save(str(path))
            result = run_lint(path, "--chapter")
            self.assertEqual(result.returncode, 0, result.stdout)
            self.assertIn("CAP-07", result.stdout)

    def test_cited_numbers_support_shorthand(self):
        """“图4.10、11、14”这种略写要能展开成三个编号。"""
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "cite.docx"
            document = docx.Document(str(FIXTURE))
            document.add_paragraph("这一解释与实验观察（图4.10、11、14）高度吻合。")
            document.add_paragraph("见表3.1~3.3。")
            document.add_paragraph("如图4.16所示。")
            document.save(str(path))
            lint = tl.Linter(path, "master", chapter_mode=True)
            cited = lint._cited_numbers()
            self.assertIn(("图", 4, 10), cited)
            self.assertIn(("图", 4, 11), cited)
            self.assertIn(("图", 4, 14), cited)
            self.assertIn(("图", 4, 16), cited)
            self.assertIn(("表", 3, 1), cited)

    def test_json_report_is_written(self):
        with tempfile.TemporaryDirectory() as td:
            report = Path(td) / "report.json"
            result = self.run_lint(FIXTURE, "--json", str(report))
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            data = json.loads(report.read_text(encoding="utf-8"))
            self.assertEqual(data["degree"], "master")
            self.assertIn("references", data["metrics"])
            self.assertIn("issues", data)


if __name__ == "__main__":
    unittest.main(verbosity=2)
