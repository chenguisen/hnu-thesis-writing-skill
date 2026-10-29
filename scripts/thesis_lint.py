#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""湖南大学研究生学位论文格式自检工具。

用法：
    python thesis_lint.py thesis.docx --degree master
    python thesis_lint.py thesis.docx --degree doctor --json report.json

只依赖 python-docx。检查页面、字体字号、标题层次、摘要/关键词、参考文献数量、
图表题注、表格边框、页眉页码等“机器能查”的项目；签名、内容质量、印刷装订需人工复核。

退出码：0 = 无 ERROR（可有 WARN）；1 = 存在 ERROR；2 = 文件无法读取。
"""

from __future__ import annotations

import argparse
import json
import re
import statistics
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Optional

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
except ImportError:  # pragma: no cover
    print("缺少 python-docx，请先安装：pip install python-docx", file=sys.stderr)
    raise SystemExit(2)

# ---------------------------------------------------------------------------
# 规范常量（2005 版）
# ---------------------------------------------------------------------------

DEGREES = {
    "doctor": {
        "label": "博士",
        "chars": (50_000, 80_000),
        "abstract": 1200,
        "refs": 80,
        "foreign_refs": 40,
    },
    "master": {
        "label": "学术硕士",
        "chars": (30_000, 50_000),
        "abstract": 600,
        "refs": 40,
        "foreign_refs": 10,
    },
    "professional": {
        "label": "专业硕士",
        "chars": (30_000, 50_000),
        "abstract": 600,
        "refs": 30,
        "foreign_refs": 5,
    },
}

SIZE = {
    "初号": 42,
    "小初": 36,
    "一号": 26,
    "小一": 24,
    "二号": 22,
    "小二": 18,
    "三号": 16,
    "小三": 15,
    "四号": 14,
    "小四": 12,
    "五号": 10.5,
    "小五": 9,
    "六号": 7.5,
}

HEITI = ("黑体", "SimHei", "Microsoft YaHei", "微软雅黑")
SONGTI = ("宋体", "SimSun", "NSimSun")
# 把西文/数字排成中文字体（校模板要求西文用 Times New Roman）
CJK_ASCII_FONTS = (
    "宋体", "黑体", "仿宋", "楷体", "微软雅黑", "华文宋体", "华文黑体",
    "SimSun", "SimHei", "NSimSun", "FangSong", "KaiTi", "Microsoft YaHei",
)
LATIN_FONT = "Times New Roman"
RE_LATIN = re.compile(r"[A-Za-z0-9]")
# 东亚宽度为“歧义”的符号：word 会按 w:hint 决定用哪套字体，全角/半角差别明显
RE_SYMBOL = re.compile(
    r"[\u00b0\u00b1\u00b5\u00b7\u00d7\u00f7"
    r"\u0370-\u03ff\u1f00-\u1fff"
    r"\u2032-\u2037\u2070-\u209f"
    r"\u2103\u2109\u210e\u2113\u212b"
    r"\u2190-\u21ff\u2200-\u22ff\u2300-\u23ff\u25a0-\u25ff]"
)

_NUM_SEP = r"(?:\s+|(?=[\u4e00-\u9fff]))"
RE_CHAPTER = re.compile(r"^第\s*(\d+)\s*章[\s　]*(.*)$")
RE_L2 = re.compile(r"^(\d+)\s*[.．]\s*(\d+)" + _NUM_SEP + r"(\S.*)$")
RE_L3 = re.compile(r"^(\d+)\s*[.．]\s*(\d+)\s*[.．]\s*(\d+)" + _NUM_SEP + r"(\S.*)$")
RE_L4 = re.compile(
    r"^(\d+)\s*[.．]\s*(\d+)\s*[.．]\s*(\d+)\s*[.．]\s*(\d+)" + _NUM_SEP + r"(\S.*)$"
)
RE_ITEM = re.compile(r"^[（(]\s*\d+\s*[）)]")
RE_REF_ENTRY = re.compile(r"^\s*[\[［]\s*(\d+)\s*[\]］]")
RE_FIG_CAP = re.compile(r"^\s*图\s*(\d+)\s*[.．]\s*(\d+)(?:\s+|$)")
RE_TAB_CAP = re.compile(r"^\s*表\s*(\d+)\s*[.．]\s*(\d+)(?:\s+|$)")
RE_FIG_LOOSE = re.compile(r"^\s*图\s*(\d+)\s*[.．]?\s+(\S.*)$")
RE_TAB_LOOSE = re.compile(r"^\s*表\s*(\d+)\s*[.．]?\s+(\S.*)$")
RE_TAB_PLACEHOLDER = re.compile(r"^\s*表\s*[xX×]\s*(?:[.．]\s*|\s+)(\S.*)$")
RE_CITATION = re.compile(r"[\[［]\s*\d+(?:\s*[-–—~～,，、]\s*\d+)*\s*[\]］]")
# 正文引用：“图4.10、11、14”“表3.1~3.3”“图4.10(a)”
RE_CITE_GROUP = re.compile(
    r"([图表])\s*(\d+)\s*[.．]\s*(\d+)((?:\s*[、,，和与及~～\-–]\s*\d+(?:[.．]\d+)?)*)"
)
RE_CJK = re.compile(r"[\u4e00-\u9fff]")
# OMML 行内公式：题注里出现时会排在文字前，并被目录/索引域一起复制
M_OMML = "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath"
M_OMML_PARA = "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMathPara"
RE_ASCII_ALPHA = re.compile(r"[A-Za-z]")

TOC_TITLES = {"目 录", "目录", "目　　录"}
ABSTRACT_TITLES = {"摘 要", "摘要", "摘　　要"}
REF_TITLES = {"参考文献", "参 考 文 献", "参　　考　　文　　献"}
BACK_TITLES = {"致 谢", "致谢", "后记"}
DECL_TITLE = "原创性声明"


# ---------------------------------------------------------------------------
# 数据模型
# ---------------------------------------------------------------------------


@dataclass
class Issue:
    level: str  # ERROR / WARN / INFO
    code: str
    message: str
    hint: str = ""


@dataclass
class Report:
    file: str
    degree: str
    issues: list[Issue] = field(default_factory=list)
    metrics: dict = field(default_factory=dict)

    def add(self, level: str, code: str, message: str, hint: str = "") -> None:
        self.issues.append(Issue(level, code, message, hint))

    def error(self, code: str, message: str, hint: str = "") -> None:
        self.add("ERROR", code, message, hint)

    def warn(self, code: str, message: str, hint: str = "") -> None:
        self.add("WARN", code, message, hint)

    def info(self, code: str, message: str, hint: str = "") -> None:
        self.add("INFO", code, message, hint)

    @property
    def errors(self) -> list[Issue]:
        return [i for i in self.issues if i.level == "ERROR"]

    @property
    def warnings(self) -> list[Issue]:
        return [i for i in self.issues if i.level == "WARN"]

    def to_dict(self) -> dict:
        return {
            "file": self.file,
            "degree": self.degree,
            "summary": {
                "errors": len(self.errors),
                "warnings": len(self.warnings),
                "infos": len(self.issues) - len(self.errors) - len(self.warnings),
            },
            "metrics": self.metrics,
            "issues": [i.__dict__ for i in self.issues],
        }


# ---------------------------------------------------------------------------
# 文档读取与格式解析
# ---------------------------------------------------------------------------


def cm(value) -> Optional[float]:
    if value is None:
        return None
    if hasattr(value, "cm"):
        return round(value.cm, 2)
    if isinstance(value, int):
        return round(value / 360000, 2)
    return None


def pt(value) -> Optional[float]:
    if value is None:
        return None
    if hasattr(value, "pt"):
        return round(value.pt, 1)
    if isinstance(value, int):
        return round(value / 12700, 1)
    return None


def _style_chain(style):
    seen = set()
    while style is not None and id(style) not in seen:
        seen.add(id(style))
        yield style
        style = style.base_style


def style_font(style, attr: str) -> Optional[str]:
    for st in _style_chain(style):
        rpr = st.element.find(qn("w:rPr"))
        if rpr is None:
            continue
        if attr in ("eastAsia", "ascii", "hAnsi"):
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is not None:
                val = fonts.get(qn("w:" + attr))
                if val:
                    return val
        elif attr == "size":
            sz = rpr.find(qn("w:sz"))
            if sz is not None:
                try:
                    return float(sz.get(qn("w:val"))) / 2.0
                except (TypeError, ValueError):
                    pass
    return None


def style_bold(style) -> Optional[bool]:
    for st in _style_chain(style):
        rpr = st.element.find(qn("w:rPr"))
        if rpr is not None:
            b = rpr.find(qn("w:b"))
            if b is not None:
                val = b.get(qn("w:val"))
                return val not in ("0", "false", "False")
    return None


def run_font(run) -> Optional[str]:
    rpr = run._element.find(qn("w:rPr"))
    if rpr is None:
        return None
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        return None
    return fonts.get(qn("w:eastAsia")) or fonts.get(qn("w:ascii"))


def run_size(run) -> Optional[float]:
    if run.font.size is not None:
        return round(run.font.size.pt, 1)
    return None


def effective_font(paragraph: Paragraph) -> Optional[str]:
    for run in paragraph.runs:
        if not run.text.strip():
            continue
        f = run_font(run)
        if f:
            return f
    return style_font(paragraph.style, "eastAsia") or style_font(paragraph.style, "ascii")


def effective_size(paragraph: Paragraph) -> Optional[float]:
    for run in paragraph.runs:
        if not run.text.strip():
            continue
        s = run_size(run)
        if s:
            return s
    return style_font(paragraph.style, "size")


def effective_bold(paragraph: Paragraph) -> Optional[bool]:
    for run in paragraph.runs:
        if run.text.strip() and run.font.bold is not None:
            return run.font.bold
    return style_bold(paragraph.style)


def style_pf_value(style, attr: str):
    """沿样式继承链读取段落格式；直接格式为 None 时使用。"""
    for st in _style_chain(style):
        val = getattr(st.paragraph_format, attr)
        if val is not None:
            return val
    return None


def effective_alignment(paragraph: Paragraph):
    al = paragraph.paragraph_format.alignment
    if al is not None:
        return al
    return style_pf_value(paragraph.style, "alignment")


def effective_first_line_indent(paragraph: Paragraph):
    val = paragraph.paragraph_format.first_line_indent
    if val is not None:
        return val
    return style_pf_value(paragraph.style, "first_line_indent")


def effective_line_spacing(paragraph: Paragraph):
    val = paragraph.paragraph_format.line_spacing
    if val is not None:
        return val
    return style_pf_value(paragraph.style, "line_spacing")


def effective_space(paragraph: Paragraph, attr: str):
    val = getattr(paragraph.paragraph_format, attr)
    if val is not None:
        return val
    return style_pf_value(paragraph.style, attr)


def paragraph_text(p: Paragraph) -> str:
    return p.text.replace("\u00a0", " ").strip()


def is_toc_style(style_name: str) -> bool:
    s = (style_name or "").lower()
    return "toc" in s or "目录" in (style_name or "")


def is_toc_line(text: str, style_name: str = "") -> bool:
    """目录/索引行：TOC 样式、点线、或“标题 + 制表符 + 页码”，不能当作标题。"""
    if is_toc_style(style_name):
        return True
    t = text.strip()
    if t.count("…") >= 2 or re.search(r"[.．·]{3,}", t):
        return True
    # Word 自动目录常见形式：标题\t页码，或标题……页码
    if "\t" in t and re.search(r"\t\s*\d+\s*$", t):
        return True
    if re.search(r"[.．·]{3,}\s*\d+\s*$", t):
        return True
    return False


def style_heading_level(style_name: str) -> Optional[int]:
    """从样式名判断是第几级标题（Heading 1 / 标题 1）。"""
    s = (style_name or "").strip().lower()
    m = re.match(r"^(?:heading|标题)\s*([1-9])$", s)
    return int(m.group(1)) if m else None


def classify_heading(text: str, style_name: str = "",
                     auto_level: Optional[int] = None) -> Optional[str]:
    t = text.strip()
    if not t or len(t) > 80:
        return None
    if is_toc_line(t, style_name):
        return None
    # 标题用 Word 自动编号（多级列表）时，正文里只有“绪论”，没有“第1章”字样
    if auto_level in (1, 2, 3, 4):
        return {1: "chapter", 2: "l2", 3: "l3", 4: "l4"}[auto_level]
    m = RE_CHAPTER.match(t)
    if m:
        title = m.group(2).strip()
        # 目录里的“第1章 绪论…………1”已由 is_toc_line 排除；标题应有中文
        if title and RE_CJK.search(title):
            return "chapter"
        return None
    bare = re.sub(r"\s+", "", t)
    if bare in {"结论", "参考文献", "致谢", "后记"}:
        return "special"
    if bare == "摘要" or bare == "目录":
        return "front"
    if re.match(r"^附录[A-Z]?(?:[：:、\s]|$)", t):
        return "appendix"
    if RE_L4.match(t):
        return "l4"
    if RE_L3.match(t):
        return "l3"
    if RE_L2.match(t):
        return "l2"
    if RE_ITEM.match(t) and len(t) <= 60:
        return "item"
    return None


def iter_blocks(doc):
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def get_section_start_type(section) -> Optional[str]:
    pg = section._sectPr.find(qn("w:pgNumType"))
    if pg is None:
        return None
    fmt = pg.get(qn("w:fmt"))
    if fmt is None and pg.get(qn("w:start")) is not None:
        # OOXML：w:pgNumType 省略 w:fmt 时按十进制处理（只设了起始页码）
        return "decimal"
    return fmt


def get_page_field_count(section) -> int:
    xml = section.footer._element.xml
    return xml.count("PAGE")


def count_ref_fields(doc) -> int:
    """统计 Word 交叉引用域（REF _Ref...）。PAGEREF 不计入。"""
    n = 0
    for el in doc.element.body.iter(qn("w:instrText")):
        text = el.text or ""
        if re.search(r"\bREF\b", text):
            n += 1
    return n


def _iter_hdr_ftr_run_elements(doc):
    try:
        from docx.parts.hdrftr import FooterPart, HeaderPart
    except ImportError:  # pragma: no cover
        return
    for part in getattr(doc.part.package, "parts", []):
        if isinstance(part, (HeaderPart, FooterPart)):
            element = getattr(part, "element", None)
            if element is not None:
                for r in element.iter(qn("w:r")):
                    yield r


def latin_in_cjk_font_usage(doc) -> tuple[int, list[str], list[str]]:
    """西文/数字被排成中文字体的 run 数、样例文字、涉及的在用样式。

    校模板（含答辩版论文）的样式是 eastAsia=宋体/黑体 + ascii=Times New Roman；
    若 ascii 被写成宋体/黑体，英文与数字就会用中文字体排。
    """
    runs = 0
    samples: list[str] = []

    def check(element):
        nonlocal runs
        for r in element.iter(qn("w:r")):
            text = "".join(t.text or "" for t in r.iter(qn("w:t")))
            if not RE_LATIN.search(text):
                continue
            rpr = r.find(qn("w:rPr"))
            rfonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
            if rfonts is None:
                continue
            if rfonts.get(qn("w:ascii")) in CJK_ASCII_FONTS:
                runs += 1
                if len(samples) < 8:
                    samples.append(text.strip()[:24])

    check(doc.element.body)
    for r in _iter_hdr_ftr_run_elements(doc):
        text = "".join(t.text or "" for t in r.iter(qn("w:t")))
        if not RE_LATIN.search(text):
            continue
        rpr = r.find(qn("w:rPr"))
        rfonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
        if rfonts is not None and rfonts.get(qn("w:ascii")) in CJK_ASCII_FONTS:
            runs += 1

    used = set()
    for p in doc.paragraphs:
        style = getattr(p, "style", None)
        if style is not None and style.name:
            used.add(style.name)
    styles: list[str] = []
    for style in doc.styles:
        if style.name not in used:
            continue
        element = getattr(style, "element", None)
        rpr = element.find(qn("w:rPr")) if element is not None else None
        rfonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
        if rfonts is not None and rfonts.get(qn("w:ascii")) in CJK_ASCII_FONTS:
            styles.append(style.name)
    return runs, samples, styles


def symbol_in_cjk_font_usage(doc) -> tuple[int, list[str], int]:
    """希腊字母/单位符号被中文字体排的 run 数、样例、涉及符号种类数。

    `w:rFonts` 带 `w:hint="eastAsia"` 时，东亚宽度为“歧义”的字符（希腊字母、′″、℃、
    °、±、×、÷、µ、下标 ₂ 等）会交给中文字体排，得到全角字形：又宽又大，与正文里
    同一符号（Times New Roman）不一致；题注、插图索引、表格里最容易出现。
    """
    runs = 0
    samples: list[str] = []
    symbols: set[str] = set()

    def check(element):
        nonlocal runs
        for r in element.iter(qn("w:r")):
            text = "".join(t.text or "" for t in r.iter(qn("w:t")))
            if not text or not RE_SYMBOL.search(text):
                continue
            rpr = r.find(qn("w:rPr"))
            rfonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
            if rfonts is None:
                continue
            if (rfonts.get(qn("w:hint")) or "").lower() != "eastasia":
                continue
            runs += 1
            symbols.update(RE_SYMBOL.findall(text))
            if len(samples) < 8:
                samples.append(text.strip()[:24])

    check(doc.element.body)
    for r in _iter_hdr_ftr_run_elements(doc):
        text = "".join(t.text or "" for t in r.iter(qn("w:t")))
        if not text or not RE_SYMBOL.search(text):
            continue
        rpr = r.find(qn("w:rPr"))
        rfonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
        if rfonts is not None and (rfonts.get(qn("w:hint")) or "").lower() == "eastasia":
            runs += 1
            symbols.update(RE_SYMBOL.findall(text))
    return runs, samples, len(symbols)


def _border_is_visible(el) -> bool:
    return el is not None and el.get(qn("w:val")) not in (None, "nil", "none")


def table_has_side_borders(table: Table) -> bool:
    # 单元格级边线优先级最高
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.tcPr
            if tc_pr is None:
                continue
            borders = tc_pr.find(qn("w:tcBorders"))
            if borders is None:
                continue
            for side in ("left", "right"):
                if _border_is_visible(borders.find(qn("w:" + side))):
                    return True

    tbl_pr = table._tbl.tblPr
    explicit: dict[str, Optional[str]] = {}
    if tbl_pr is not None:
        borders = tbl_pr.find(qn("w:tblBorders"))
        if borders is not None:
            for side in ("left", "right"):
                el = borders.find(qn("w:" + side))
                if el is not None:
                    explicit[side] = el.get(qn("w:val"))

    # 显式设置优先于表格样式：左右都设为 nil 时不应再按 Table Grid 误报
    if any(_border_is_visible_val(v) for v in explicit.values()):
        return True
    if len(explicit) == 2 and all(not _border_is_visible_val(v) for v in explicit.values()):
        return False

    style_name = getattr(getattr(table, "style", None), "name", "") or ""
    return style_name in ("Table Grid", "网格型", "网格表")


def _border_is_visible_val(val: Optional[str]) -> bool:
    return val not in (None, "nil", "none")


DOUBLE_LINE_VALUES = {
    "double",
    "thinThickSmallGap",
    "thinThickMediumGap",
    "thinThickLargeGap",
    "thickThinSmallGap",
    "thickThinMediumGap",
    "thickThinLargeGap",
}


def header_has_double_border(section) -> bool:
    for hdr in (section.header, section.even_page_header):
        try:
            paras = hdr.paragraphs
        except Exception:
            continue
        for p in paras:
            ppr = p._p.pPr
            if ppr is None:
                continue
            pbdr = ppr.find(qn("w:pBdr"))
            if pbdr is not None:
                bottom = pbdr.find(qn("w:bottom"))
                if bottom is not None and bottom.get(qn("w:val")) in DOUBLE_LINE_VALUES:
                    return True
    return False


# ---------------------------------------------------------------------------
# 检查器
# ---------------------------------------------------------------------------


class Linter:
    def __init__(self, path: Path, degree: str, chapter_mode: bool = False):
        self.path = path
        self.degree = degree
        self.chapter_mode = chapter_mode
        self.cfg = DEGREES[degree]
        self.doc = Document(str(path))
        self.paras = [p for p in self.doc.paragraphs]
        self.blocks = list(iter_blocks(self.doc))
        self.texts = [paragraph_text(p) for p in self.paras]
        self.report = Report(str(path), degree)
        self._body_fonts: list[tuple[str, float]] = []
        self._line_spacings: list[float] = []

    # -- 标题识别 ---------------------------------------------------------
    def _numbering_style_levels(self) -> dict[str, int]:
        """numbering.xml 里“多级列表链接到样式”的情况：styleId -> 标题级别。"""
        if getattr(self, "_num_style_levels", None) is not None:
            return self._num_style_levels
        levels: dict[str, int] = {}
        try:
            numbering = self.doc.part.numbering_part.element
        except Exception:  # pragma: no cover - 文档没有 numbering 部件
            self._num_style_levels = levels
            return levels
        abs_map: dict[str, dict[str, int]] = {}
        for an in numbering.findall(qn("w:abstractNum")):
            aid = an.get(qn("w:abstractNumId"))
            for lvl in an.findall(qn("w:lvl")):
                pstyle = lvl.find(qn("w:pStyle"))
                if pstyle is None:
                    continue
                try:
                    level = int(lvl.get(qn("w:ilvl"), "0")) + 1
                except ValueError:
                    level = 1
                abs_map.setdefault(aid, {})[pstyle.get(qn("w:val"))] = level
        for num in numbering.findall(qn("w:num")):
            aid_el = num.find(qn("w:abstractNumId"))
            if aid_el is None:
                continue
            for style_id, level in abs_map.get(aid_el.get(qn("w:val")), {}).items():
                levels.setdefault(style_id, level)
        self._num_style_levels = levels
        return levels

    def _heading_auto_level(self, p: Paragraph) -> Optional[int]:
        """段落是带自动编号（numId≠0）的第几级标题；不是则返回 None。"""
        level = style_heading_level(p.style.name)
        if level is None:
            return None
        num_id = None
        ppr = p._p.pPr
        if ppr is not None:
            numpr = ppr.find(qn("w:numPr"))
            if numpr is not None:
                nid = numpr.find(qn("w:numId"))
                if nid is not None:
                    num_id = nid.get(qn("w:val"))
        if num_id is None:
            style = p.style
            while style is not None:
                sppr = style.element.pPr
                if sppr is not None:
                    numpr = sppr.find(qn("w:numPr"))
                    if numpr is not None:
                        nid = numpr.find(qn("w:numId"))
                        if nid is not None:
                            num_id = nid.get(qn("w:val"))
                            break
                style = style.base_style
        if num_id in (None, "0"):
            if num_id == "0":
                return None
            # 多级列表直接“链接到样式”时，段落和样式上都没有 numPr
            return self._numbering_style_levels().get(p.style.style_id)
        return level

    def _classify(self, p: Paragraph, text: str) -> Optional[str]:
        return classify_heading(text, p.style.name,
                                auto_level=self._heading_auto_level(p))

    # -- 页面 -------------------------------------------------------------
    def check_pages(self) -> None:
        if not self.doc.sections:
            self.report.error("PAGE-01", "文档没有节（section）信息，无法检查页面设置")
            return
        for i, s in enumerate(self.doc.sections, 1):
            w, h = cm(s.page_width), cm(s.page_height)
            if w is None or h is None:
                continue
            if abs(w - 21.0) > 0.2 or abs(h - 29.7) > 0.2:
                self.report.error(
                    "PAGE-02",
                    f"第 {i} 节纸张为 {w}×{h} cm，不是 A4（21.0×29.7 cm）",
                    "页面设置 → 纸张 → A4",
                )
            margins = {
                "上": cm(s.top_margin),
                "下": cm(s.bottom_margin),
                "左": cm(s.left_margin),
                "右": cm(s.right_margin),
            }
            expected = {"上": 2.2, "下": 2.2, "左": 2.5, "右": 2.5}
            for name, val in margins.items():
                if val is None:
                    continue
                if abs(val - expected[name]) > 0.12:
                    self.report.warn(
                        "PAGE-03",
                        f"第 {i} 节{name}页边距为 {val} cm，规范为 {expected[name]} cm",
                        "若学院模板不同，以模板为准；版心 155 mm 时左侧可留 3.0 cm 装订线",
                    )
            for name, val, exp in (
                ("页眉", cm(s.header_distance), 1.1),
                ("页脚", cm(s.footer_distance), 1.1),
            ):
                if val is not None and abs(val - exp) > 0.12:
                    self.report.warn(
                        "PAGE-04",
                        f"第 {i} 节{name}距边界为 {val} cm，规范为 {exp} cm",
                    )

    # -- 页眉 / 页码 ------------------------------------------------------
    def check_headers(self) -> None:
        if not self.doc.sections:
            return
        if self.chapter_mode:
            return
        body_section = self.doc.sections[-1]
        header_text = "\n".join(p.text for p in body_section.header.paragraphs).strip()
        if not header_text:
            self.report.warn(
                "HEAD-01",
                "正文节的页眉为空",
                "页眉应为奇数页“学位论文类别”、偶数页“论文题目”",
            )
        even_text = "\n".join(p.text for p in body_section.even_page_header.paragraphs).strip()
        if not even_text:
            self.report.warn(
                "HEAD-02",
                "未检测到偶数页页眉内容",
                "页面设置 → 版式 → 勾选“奇偶页不同”，分别填写奇偶页眉",
            )
        if not header_has_double_border(body_section):
            self.report.warn(
                "HEAD-03",
                "页眉段落未检测到双线边框",
                "页眉下边框设为双线、1.5 磅，粗线在上",
            )
        numbered_sections = self.doc.sections[1:] if len(self.doc.sections) > 1 else self.doc.sections
        first_number = 2 if len(self.doc.sections) > 1 else 1
        missing_pages = [
            i
            for i, s in enumerate(numbered_sections, start=first_number)
            if get_page_field_count(s) == 0
        ]
        if missing_pages:
            self.report.warn(
                "PAGE-05",
                f"第 {missing_pages} 节页脚未检测到 PAGE 域",
                "插入页码域，不要手工输入数字；前置罗马、正文阿拉伯需分节设置",
            )
        fmts = [get_section_start_type(s) for s in self.doc.sections]
        has_roman = any(f in ("lowerRoman", "upperRoman") for f in fmts)
        has_decimal = any(f == "decimal" for f in fmts)
        if len(self.doc.sections) > 1 and not has_roman:
            self.report.warn(
                "PAGE-06",
                "未检测到罗马数字页码格式",
                "前置部分（原创性声明起）应为罗马数字，正文为阿拉伯数字",
            )
        if len(self.doc.sections) > 1 and not has_decimal:
            self.report.warn("PAGE-07", "未检测到阿拉伯数字页码格式")

    # -- 标题 -------------------------------------------------------------
    def check_headings(self) -> None:
        counts = {"chapter": 0, "special": 0, "l2": 0, "l3": 0, "l4": 0, "item": 0, "appendix": 0}
        current_chapter = 0
        level2_seq: dict[int, list[int]] = {}
        level3_seq: dict[tuple[int, int], list[int]] = {}
        chapter_numbers: list[int] = []

        for p, t in zip(self.paras, self.texts):
            if not t:
                continue
            kind = self._classify(p, t)
            if kind is None:
                continue
            if kind in counts:
                counts[kind] += 1
            font = effective_font(p) or ""
            size = effective_size(p)
            bold = effective_bold(p)

            if kind in ("chapter", "special"):
                # “结论/参考文献/致谢”也可能是论文构成清单中的一行；
                # 只有大字号或居中时才按章标题检查，避免把目录清单当正文。
                if kind == "special":
                    is_heading = (size is not None and size >= 16) or effective_alignment(p) == WD_ALIGN_PARAGRAPH.CENTER
                    if not is_heading:
                        counts["special"] -= 1
                        continue
                m = RE_CHAPTER.match(t)
                if m:
                    num = int(m.group(1))
                    if not chapter_numbers:
                        chapter_numbers.append(num)
                    else:
                        expected = chapter_numbers[0] + len(chapter_numbers)
                        if num != expected:
                            self.report.warn(
                                "HEAD-04",
                                f"章编号不连续：期望第 {expected} 章，实际“{t[:20]}”",
                            )
                        chapter_numbers.append(num)
                    current_chapter = num
                elif kind == "chapter":
                    # 自动编号章：正文中没有章号文字，按出现顺序推算
                    expected = chapter_numbers[0] + len(chapter_numbers) if chapter_numbers else 1
                    chapter_numbers.append(expected)
                    current_chapter = expected
                self._check_heading_style(p, t, "章", 18, "center", font, size, bold)
                before = effective_space(p, "space_before")
                after = effective_space(p, "space_after")
                if before is not None and before.pt < 12:
                    self.report.warn("HEAD-05", f"“{t[:20]}”段前间距偏小，章标题上下应各空一行")
                if after is not None and after.pt < 12:
                    self.report.warn("HEAD-06", f"“{t[:20]}”段后间距偏小，章标题上下应各空一行")

            elif kind == "l2":
                m = RE_L2.match(t)
                if m:
                    ch, sec = int(m.group(1)), int(m.group(2))
                    if ch != current_chapter:
                        if self.chapter_mode and current_chapter == 0:
                            # 单章草稿：没有“第3章”标题时，以第一个 3.x 的章号为准
                            current_chapter = ch
                        else:
                            self.report.warn(
                                "HEAD-07",
                                f"“{t[:24]}”的章号 {ch} 与当前章 {current_chapter} 不一致",
                            )
                    level2_seq.setdefault(ch, []).append(sec)
                self._check_heading_style(p, t, "节（1.1）", 15, "left", font, size, bold)
                self._check_space(p, t, 12)

            elif kind == "l3":
                m = RE_L3.match(t)
                if m:
                    ch, sec, sub = int(m.group(1)), int(m.group(2)), int(m.group(3))
                    key = (ch, sec)
                    level3_seq.setdefault(key, []).append(sub)
                self._check_heading_style(p, t, "条（1.1.1）", 14, "left", font, size, bold)
                self._check_space(p, t, 6)

            elif kind == "l4":
                self._check_heading_style(p, t, "款（1.1.1.1）", 12, "left", font, size, bold)
                self._check_space(p, t, 6)

            elif kind == "appendix":
                self._check_heading_style(p, t, "附录标题", 18, "center", font, size, bold, strict=False)

        self.report.metrics["headings"] = counts
        if counts["chapter"] == 0:
            if self.chapter_mode and counts["l2"] > 0:
                self.report.warn(
                    "HEAD-15",
                    "单章模式：未检测到“第X章”标题；请确认章号，并把章标题设为小 2 号黑体居中",
                )
            else:
                self.report.error("HEAD-08", "未识别到任何章标题（“第X章 …”）")
        if counts["l2"] == 0:
            self.report.warn("HEAD-09", "未识别到节标题（“1.1 …”）")

        # 跳号检查
        for ch, nums in sorted(level2_seq.items()):
            if nums != list(range(1, len(nums) + 1)):
                self.report.warn("HEAD-10", f"第 {ch} 章节编号可能有跳号：{nums}")
        for (ch, sec), nums in sorted(level3_seq.items()):
            if nums != list(range(1, len(nums) + 1)):
                self.report.warn("HEAD-10", f"第 {ch}.{sec} 条编号可能有跳号：{nums}")

    def _check_heading_style(
        self,
        p: Paragraph,
        text: str,
        label: str,
        size: float,
        align: str,
        font: str,
        actual_size: Optional[float],
        bold: Optional[bool],
        strict: bool = True,
    ) -> None:
        if actual_size is not None and abs(actual_size - size) > 0.6:
            self.report.warn(
                "FONT-01",
                f"{label}“{text[:20]}”字号为 {actual_size} pt，规范为 {size} pt",
                f"在样式中设置，不要逐段手工改",
            )
        if strict and font and not any(f in font for f in HEITI):
            self.report.warn(
                "FONT-02",
                f"{label}“{text[:20]}”字体为 {font}，规范为黑体",
            )
        actual_align = effective_alignment(p)
        if align == "center" and actual_align not in (WD_ALIGN_PARAGRAPH.CENTER,):
            self.report.warn(
                "HEAD-11",
                f"{label}“{text[:20]}”未居中",
            )
        if align == "left" and actual_align == WD_ALIGN_PARAGRAPH.CENTER:
            self.report.warn(
                "HEAD-12",
                f"{label}“{text[:20]}”应顶格左对齐",
            )

    def _check_space(self, p: Paragraph, text: str, expected: float) -> None:
        before = effective_space(p, "space_before")
        after = effective_space(p, "space_after")
        if before is not None and before.pt < expected - 1:
            self.report.warn(
                "HEAD-13",
                f"“{text[:20]}”段前 {before.pt:.0f} 磅，规范为 {expected:.0f} 磅",
            )
        if after is not None and after.pt < expected - 1:
            self.report.warn(
                "HEAD-14",
                f"“{text[:20]}”段后 {after.pt:.0f} 磅，规范为 {expected:.0f} 磅",
            )

    # -- 正文 -------------------------------------------------------------
    def check_body(self) -> None:
        start = None
        end = len(self.texts)
        for i, t in enumerate(self.texts):
            if start is None and self._classify(self.paras[i], t) == "chapter":
                start = i
            if start is not None and re.sub(r"\s+", "", t) in {"参考文献"}:
                end = i
                break
        if start is None:
            if self.chapter_mode:
                for i, t in enumerate(self.texts):
                    if self._classify(self.paras[i], t) in ("l2", "l3"):
                        start = i
                        break
            if start is None:
                return

        sizes: list[float] = []
        fonts: list[str] = []
        indents = 0
        body_count = 0
        multiple_spacings: list[float] = []
        for p, t in zip(self.paras[start:end], self.texts[start:end]):
            if not t or self._classify(p, t) is not None:
                continue
            if RE_FIG_CAP.match(t) or RE_TAB_CAP.match(t):
                continue
            if len(RE_CJK.findall(t)) < 10:
                continue
            body_count += 1
            s = effective_size(p)
            f = effective_font(p)
            if s:
                sizes.append(s)
            if f:
                fonts.append(f)
            indent = effective_first_line_indent(p)
            if indent is not None and indent.cm >= 0.5:
                indents += 1
            ls = effective_line_spacing(p)
            if ls is not None:
                if hasattr(ls, "pt"):
                    self._line_spacings.append(ls.pt)
                elif isinstance(ls, (int, float)):
                    multiple_spacings.append(float(ls))

        if sizes:
            main_size = statistics.mode(sizes)
            self.report.metrics["body_size_pt"] = main_size
            if abs(main_size - 12.0) > 0.6:
                self.report.warn(
                    "FONT-03",
                    f"正文主字号为 {main_size} pt，规范为小 4 号（12 pt）",
                )
        if fonts:
            main_font = statistics.mode(fonts)
            self.report.metrics["body_font"] = main_font
            if not any(f in main_font for f in SONGTI):
                self.report.warn(
                    "FONT-04",
                    f"正文主字体为 {main_font}，规范为宋体",
                )
        if body_count and indents / body_count < 0.5:
            self.report.warn(
                "FONT-05",
                "超过一半的正文段落没有首行缩进 2 字符",
                "用“首行缩进 2 字符”或样式统一设置",
            )
        if self._line_spacings:
            med = statistics.median(self._line_spacings)
            self.report.metrics["body_line_spacing_pt"] = round(med, 1)
            if not (18 <= med <= 24):
                self.report.warn(
                    "FONT-06",
                    f"正文固定行距中位数为 {med} pt，通常应为 20 磅左右（每页 35 行）",
                )
        elif body_count:
            if multiple_spacings:
                med = statistics.median(multiple_spacings)
                self.report.warn(
                    "FONT-06",
                    f"正文未设置固定行距，{len(multiple_spacings)} 段为 {med:g} 倍行距；"
                    "建议改为固定 20 磅，按每页 35 行控制",
                )
            else:
                self.report.warn(
                    "FONT-06",
                    "正文未设置固定行距（当前为默认单倍行距）；建议改为固定 20 磅，按每页 35 行控制",
                )

        chars = sum(len(RE_CJK.findall(t)) for t in self.texts[start:end])
        self.report.metrics["body_chinese_chars"] = chars
        if self.chapter_mode:
            return
        lo, hi = self.cfg["chars"]
        if chars < lo * 0.9:
            self.report.warn(
                "BODY-01",
                f"正文汉字约 {chars} 字，低于{self.cfg['label']}要求（{lo}～{hi} 字）",
            )
        elif chars > hi * 1.1:
            self.report.warn(
                "BODY-02",
                f"正文汉字约 {chars} 字，高于{self.cfg['label']}要求（{lo}～{hi} 字）",
            )

        # 正文引用标注
        body_text = "\n".join(self.texts[start:end])
        cites = len(RE_CITATION.findall(body_text))
        self.report.metrics["in_text_citations"] = cites
        if cites == 0:
            self.report.warn("CITE-01", "正文未检测到 [n] 形式的引用标注")

    # -- 摘要 / 关键词 ----------------------------------------------------
    def check_abstract(self) -> None:
        if self.chapter_mode:
            return
        abs_idx = None
        for i, t in enumerate(self.texts):
            if re.sub(r"\s+", "", t) in {"摘要", "摘要：", "摘要:"}:
                abs_idx = i
                break
            if t in ABSTRACT_TITLES:
                abs_idx = i
                break
        if abs_idx is None:
            self.report.error("ABS-01", "未找到中文摘要标题“摘 要”")
        else:
            end = len(self.texts)
            for j in range(abs_idx + 1, len(self.texts)):
                tt = re.sub(r"\s+", "", self.texts[j])
                if tt.startswith("关键词") or tt.lower().startswith("abstract") or tt in {"目录"}:
                    end = j
                    break
            content = "\n".join(self.texts[abs_idx + 1:end])
            n = len(RE_CJK.findall(content))
            self.report.metrics["abstract_chars"] = n
            target = self.cfg["abstract"]
            if n < target * 0.55:
                self.report.warn(
                    "ABS-02",
                    f"中文摘要约 {n} 字，低于“约 {target} 字”的要求",
                )
            elif n > target * 1.6:
                self.report.warn(
                    "ABS-03",
                    f"中文摘要约 {n} 字，明显超过“约 {target} 字”的要求",
                )
            self.report.metrics["chinese_abstract"] = n

        kw = None
        for t in self.texts:
            if re.sub(r"\s+", "", t).startswith("关键词"):
                kw = t
                break
        if kw is None:
            self.report.warn("ABS-04", "未找到“关键词：”段落")
        else:
            body = re.sub(r"^关键词\s*[：:]\s*", "", kw)
            parts = [x for x in re.split(r"[；;、,，]", body) if x.strip()]
            self.report.metrics["keywords"] = len(parts)
            if not (3 <= len(parts) <= 8):
                self.report.warn(
                    "ABS-05",
                    f"关键词 {len(parts)} 个，规范为 3～8 个",
                )

        if not any(t.strip().lower() == "abstract" for t in self.texts):
            self.report.warn("ABS-06", "未找到英文摘要标题“Abstract”")
        if not any(re.sub(r"\s+", "", t).lower().startswith("keywords") for t in self.texts):
            self.report.warn("ABS-07", "未找到“Key Words:”段落")

    # -- 参考文献 ---------------------------------------------------------
    def check_references(self) -> None:
        if self.chapter_mode:
            return
        ref_idx = None
        for i, t in enumerate(self.texts):
            if re.sub(r"\s+", "", t) in {"参考文献"}:
                ref_idx = i
                break
        if ref_idx is None:
            self.report.error("REF-01", "未找到“参考文献”章节")
            return

        entries: list[tuple[int, str]] = []
        for t in self.texts[ref_idx + 1:]:
            if re.sub(r"\s+", "", t) in {"致谢", "附录"} or t.startswith("附录"):
                if entries:
                    break
            m = RE_REF_ENTRY.match(t)
            if m:
                entries.append((int(m.group(1)), t))
        self.report.metrics["references"] = len(entries)
        if not entries:
            self.report.error("REF-02", "参考文献章节未识别到 [1] … 形式的条目")
            return

        numbers = [n for n, _ in entries]
        expected = list(range(1, len(entries) + 1))
        if numbers != expected:
            self.report.warn(
                "REF-03",
                f"参考文献编号不连续：{numbers[:8]}…（共 {len(entries)} 条）",
                "按正文首次出现顺序从 [1] 连续编号",
            )

        foreign = 0
        for _, text in entries:
            body = RE_REF_ENTRY.sub("", text)
            ascii_n = len(RE_ASCII_ALPHA.findall(body))
            cjk_n = len(RE_CJK.findall(body))
            if ascii_n >= 20 and cjk_n <= 8:
                foreign += 1
        self.report.metrics["foreign_references"] = foreign
        if len(entries) < self.cfg["refs"]:
            self.report.warn(
                "REF-04",
                f"参考文献 {len(entries)} 条，{self.cfg['label']}要求 ≥{self.cfg['refs']} 条",
            )
        if foreign < self.cfg["foreign_refs"]:
            self.report.warn(
                "REF-05",
                f"外文文献约 {foreign} 条，{self.cfg['label']}要求 ≥{self.cfg['foreign_refs']} 条",
                "按条目语言人工复核；中文文献中的英文标题不算外文文献",
            )

    # -- 图表 -------------------------------------------------------------
    def check_figures_tables(self) -> None:
        fig_nums: dict[int, list[int]] = {}
        tab_nums: dict[int, list[int]] = {}
        loose_figs: list[int] = []
        loose_tabs: list[int] = []
        table_placeholders: list[str] = []
        for p, t in zip(self.paras, self.texts):
            if is_toc_line(t, p.style.name):
                continue
            m = RE_FIG_CAP.match(t)
            if m:
                fig_nums.setdefault(int(m.group(1)), []).append(int(m.group(2)))
            else:
                m = RE_FIG_LOOSE.match(t)
                if m:
                    loose_figs.append(int(m.group(1)))
            m = RE_TAB_CAP.match(t)
            if m:
                tab_nums.setdefault(int(m.group(1)), []).append(int(m.group(2)))
            else:
                m = RE_TAB_LOOSE.match(t)
                if m:
                    loose_tabs.append(int(m.group(1)))
            if RE_TAB_PLACEHOLDER.match(t):
                table_placeholders.append(t[:24])

        def check_seq(kind: str, data: dict[int, list[int]]) -> int:
            total = 0
            for ch, nums in sorted(data.items()):
                total += len(nums)
                if nums != list(range(1, len(nums) + 1)):
                    self.report.warn(
                        "CAP-01",
                        f"第 {ch} 章{kind}编号不连续：{nums}",
                        f"{kind}按章从 1 开始编号",
                    )
            return total

        n_fig = check_seq("图", fig_nums)
        n_tab = check_seq("表", tab_nums)
        if loose_figs:
            shown = "、".join(f"图{n}" for n in loose_figs[:5])
            self.report.warn(
                "CAP-03",
                f"{len(loose_figs)} 个图题未按章编号（{shown}…），规范为“图<章号>.<序号>”（如 图3.1）",
                "用题注并勾选“包含章节号”，正文用交叉引用；不要用 图1、图2 的全文连续编号",
            )
            n_fig += len(loose_figs)
        if loose_tabs:
            shown = "、".join(f"表{n}" for n in loose_tabs[:5])
            self.report.warn(
                "CAP-03",
                f"{len(loose_tabs)} 个表题未按章编号（{shown}…），规范为“表<章号>.<序号>”（如 表3.1）",
            )
            n_tab += len(loose_tabs)
        if table_placeholders:
            self.report.warn(
                "CAP-04",
                f"检测到 {len(table_placeholders)} 个表题仍是占位编号“表x”：{table_placeholders[0]}",
                "定稿前填入真实表号（如表3.1），并检查正文“表x”引用",
            )
            n_tab += len(table_placeholders)
        self.report.metrics["figures"] = n_fig
        self.report.metrics["tables"] = n_tab

        # 表题应在表格上方；图题应在图下方。按文档块顺序做近似检查。
        block_list = list(iter_blocks(self.doc))
        for i, block in enumerate(block_list):
            if not isinstance(block, Paragraph):
                continue
            t = paragraph_text(block)
            if is_toc_line(t, block.style.name):
                continue
            if RE_TAB_CAP.match(t):
                nxt = block_list[i + 1] if i + 1 < len(block_list) else None
                if not isinstance(nxt, Table):
                    self.report.warn(
                        "CAP-02",
                        f"表题“{t[:24]}”下方未紧跟表格（表题应在表上）",
                    )
            if RE_FIG_CAP.match(t):
                prev = block_list[i - 1] if i > 0 else None
                if isinstance(prev, Paragraph) and not prev._p.findall(".//" + qn("w:drawing")):
                    # 图题可能紧跟空段落；只提示，不判错
                    pass

        for ti, table in enumerate(self.doc.tables, 1):
            if table_has_side_borders(table):
                self.report.warn(
                    "TAB-01",
                    f"第 {ti} 个表格带有左/右边线，规范要求表格不加左右边线",
                )
            for row in table.rows:
                for cell in row.cells:
                    if "同上" in cell.text or "〃" in cell.text:
                        self.report.warn(
                            "TAB-02",
                            f"第 {ti} 个表格出现“同上/〃”，空白格应填“—”",
                        )
                        break

    # -- 交叉引用 ---------------------------------------------------------
    def check_cross_references(self) -> None:
        ref_fields = count_ref_fields(self.doc)
        self.report.metrics["ref_fields"] = ref_fields

        captions: list[tuple[str, int, int, int, bool]] = []
        for i, (p, t) in enumerate(zip(self.paras, self.texts)):
            if is_toc_line(t, p.style.name):
                continue
            m = RE_FIG_CAP.match(t)
            if m:
                captions.append(("图", int(m.group(1)), int(m.group(2)), i, False))
                continue
            m = RE_TAB_CAP.match(t)
            if m:
                captions.append(("表", int(m.group(1)), int(m.group(2)), i, False))
                continue
            m = RE_FIG_LOOSE.match(t)
            if m:
                captions.append(("图", 0, int(m.group(1)), i, True))
                continue
            m = RE_TAB_LOOSE.match(t)
            if m:
                captions.append(("表", 0, int(m.group(1)), i, True))

        has_math = bool(self.doc.element.body.findall(".//" + qn("m:oMath")))
        if (captions or has_math) and ref_fields == 0:
            what = f"{len(captions)} 个图/表题注" + ("和公式" if has_math else "")
            self.report.warn(
                "XREF-01",
                f"检测到 {what}，但未检测到 REF 交叉引用域；正文编号可能是手打，插入/删除后不会自动更新，也不能点击跳转",
                "用“插入题注 + 交叉引用（勾选插入为超链接）”；公式用书签 + 交叉引用，详见 references/cross-references.md",
            )

        orphans: list[str] = []
        cited = self._cited_numbers({c[3] for c in captions})
        for kind, ch, num, idx, loose in captions:
            if loose:
                label = f"{kind}{num}"
            else:
                label = f"{kind}{ch}.{num}"
            if (kind, ch, num) not in cited:
                orphans.append(label)
        if orphans:
            self.report.warn(
                "XREF-02",
                "以下图/表未在正文中找到引用：" + "、".join(orphans),
                "每张图/每个表都应在正文中至少引用一次；编号建议用交叉引用域",
            )

    def _cited_numbers(self, exclude: set[int] | None = None) -> set[tuple[str, int, int]]:
        """正文里引用过的图/表编号集合，支持“图4.10、11、14”“表3.1~3.3”这类略写。

        题注自身不算“被正文引用”，因此通过 exclude 跳过题注段落。
        """
        exclude = exclude or set()
        cited: set[tuple[str, int, int]] = set()
        for j, t in enumerate(self.texts):
            if j in exclude:
                continue
            if is_toc_line(t, self.paras[j].style.name):
                continue
            for m in RE_CITE_GROUP.finditer(t):
                kind, ch, num, tail = m.group(1), int(m.group(2)), int(m.group(3)), m.group(4) or ""
                cited.add((kind, ch, num))
                for part in re.split(r"[\s、,，和与及~～\-–]+", tail):
                    part = part.strip()
                    if not part:
                        continue
                    if re.fullmatch(r"\d+[.．]\d+", part):          # 4.14
                        c2, n2 = re.split(r"[.．]", part)
                        cited.add((kind, int(c2), int(n2)))
                    elif part.isdigit():                            # 11
                        cited.add((kind, ch, int(part)))
        return cited

    # -- 前置部分 ---------------------------------------------------------
    def check_front_matter(self) -> None:
        if self.chapter_mode:
            return
        all_text = "\n".join(self.texts)
        if DECL_TITLE not in all_text:
            self.report.error("FRONT-01", "未找到“学位论文原创性声明”")
        if "版权使用授权书" not in all_text:
            self.report.error("FRONT-02", "未找到“学位论文版权使用授权书”")
        if not any(re.sub(r"\s+", "", t) in {"目录"} for t in self.texts):
            self.report.warn("FRONT-03", "未找到目录")
        if "湖南大学" not in all_text:
            self.report.warn("FRONT-04", "全文未出现“湖南大学”，请确认封面/内封")

    # -- 中英文字体分工 ---------------------------------------------------
    def check_latin_fonts(self) -> None:
        runs, samples, styles = latin_in_cjk_font_usage(self.doc)
        self.report.metrics["latin_cjk_runs"] = runs
        if runs < 5 and not styles:
            return
        parts = [f"检测到 {runs} 处西文/数字用中文字体排"]
        if styles:
            shown = "、".join(styles[:4]) + ("等" if len(styles) > 4 else "")
            parts.append(f"样式 {shown} 的西文字体也是中文字体")
        if samples:
            parts.append("如“" + "”“".join(samples[:3]) + "”")
        self.report.warn(
            "FONT-07",
            "；".join(parts) + f"。校模板要求西文与数字用 {LATIN_FONT}（中文字体不变）",
            "python scripts/fix_latin_fonts.py thesis.docx --in-place；"
            f"或在 Word 里把样式/文字的“西文字体”设为 {LATIN_FONT}，中文字体保持宋体/黑体",
        )

    # -- 符号被中文字体排 ---------------------------------------------
    def check_symbol_fonts(self) -> None:
        runs, samples, kinds = symbol_in_cjk_font_usage(self.doc)
        self.report.metrics["symbol_cjk_runs"] = runs
        if runs < 3:
            return
        parts = [f"检测到 {runs} 处希腊字母/单位符号用了中文字体（{kinds} 种符号）"]
        if samples:
            parts.append("如“" + "”“".join(samples[:3]) + "”")
        self.report.warn(
            "FONT-09",
            "；".join(parts) + "。同一符号在正文里是 "
            f"{LATIN_FONT}、在题注/索引/表格里却是宋体或黑体的全角字形，"
            "宽度和粗细都对不上（通常是因为 run 带 w:hint=\"eastAsia\"）",
            "python scripts/fix_symbol_fonts.py thesis.docx --in-place；"
            "或在 Word 里选中该符号，把“西文字体”设为 "
            f"{LATIN_FONT}、并把“东亚字体提示”去掉（见 references/formatting.md 第 2.1 节）",
        )

    # -- 题注版式：不跨页、字号随样式、别夹公式 ---------------------------
    def check_caption_layout(self) -> None:
        caps = []
        for p, t in zip(self.paras, self.texts):
            low = (p.style.name or "").lower()
            if low not in ("caption", "题注"):
                continue
            if not (RE_FIG_CAP.match(t) or RE_TAB_CAP.match(t)):
                continue
            caps.append((p, t))
        if not caps:
            return

        def flag(p, tag):
            """段落属性里的开关是否打开（缺省视为关）。"""
            ppr = p._element.pPr
            if ppr is None:
                return False
            el = ppr.find(qn("w:" + tag))
            if el is None:
                return False
            return (el.get(qn("w:val")) or "true").lower() not in ("0", "false", "off")

        # ① 题注里夹了行内公式对象（会跟着复制进插图和附表索引）
        math_caps = [t for p, t in caps
                     if p._element.findall(M_OMML) or p._element.findall(M_OMML_PARA)]
        self.report.metrics["caption_math"] = len(math_caps)
        if math_caps:
            self.report.warn(
                "CAP-05",
                f"{len(math_caps)} 条题注里夹着行内公式对象（如“{math_caps[0][:18]}”）。"
                "公式会排在题注文字前面，并被目录/索引域一起复制过去",
                "把题注里的公式删掉（晶面指数等直接打“(10-10)”），"
                "或用纯文本重打；删完更新一次目录与插图索引",
            )

        # ② 题注没有 keepLines / 图题没有 keepNext：会被拆到两页
        split_risk = []
        for p, t in caps:
            need_lines = not flag(p, "keepLines")
            need_next = RE_FIG_CAP.match(t) is not None and not flag(p, "keepNext")
            if need_lines or need_next:
                split_risk.append(t[:18])
        if split_risk:
            self.report.warn(
                "CAP-06",
                f"{len(split_risk)} 条题注没设“段中不分页/与下段同页”（如“{split_risk[0]}”），"
                "换页时图题可能被拆到两页，违反“插图与其图题为一个整体”",
                "题注段落勾选“段中不分页”，图题的“与下段同页”也勾上；"
                "图所在段落勾“与下段同页”，让图与图题一起走（见 references/formatting.md 第 8 节）",
            )

        # ③ 题注字号被 run 直格式覆盖（样式是五号，run 写死 12 pt 很常见）
        style_size = None
        style_name = None
        s = caps[0][0].style
        while s is not None:
            rpr = s.element.find(qn("w:rPr"))
            sz = rpr.find(qn("w:sz")) if rpr is not None else None
            if sz is not None:
                style_size = int(sz.get(qn("w:val"))) / 2
                style_name = s.name
                break
            s = s.base_style
        overridden = []
        for p, t in caps:
            for r in p.runs:
                if not (r.text or "").strip():
                    continue
                if r.font.size is not None and style_size is not None \
                        and abs(r.font.size.pt - style_size) > 0.1:
                    overridden.append((t[:18], r.font.size.pt))
                    break
        self.report.metrics["caption_size_override"] = len(overridden)
        if overridden:
            shown = overridden[0]
            self.report.warn(
                "FONT-10",
                f"{len(overridden)} 条题注的字号被 run 直接改成 {shown[1]:g} pt，"
                f"覆盖了“{style_name}”样式的 {style_size:g} pt（如“{shown[0]}”）。"
                "题注应为五号（10.5 pt）",
                "选中题注 → 开始 → 字号恢复为样式值（Ctrl+Space 清除直格式），"
                "或 python -c 里的 python-docx 一行：删除题注 run 的 w:sz / w:szCs",
            )

    # -- 正文里不该出现的黑体 ---------------------------------------------
    def check_body_heiti(self) -> None:
        """正文用宋体；正文里出现黑体，多半是交叉引用继承了题注的格式。"""
        start = None
        for i, t in enumerate(self.texts):
            if self._classify(self.paras[i], t) == "chapter":
                start = i
                break
        if start is None:
            return
        bad = 0
        samples: list[str] = []
        for p, t in zip(self.paras[start:], self.texts[start:]):
            if self._classify(p, t) is not None:
                continue
            style_name = p.style.name if p.style is not None else ""
            low = style_name.lower()
            if low.startswith("toc") or low.startswith("caption") or "题注" in style_name:
                continue
            for run in p.runs:
                text = (run.text or "").strip()
                if not text:
                    continue
                rpr = run._element.find(qn("w:rPr"))
                rfonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
                east = rfonts.get(qn("w:eastAsia")) if rfonts is not None else None
                if east is None:
                    east = style_font(p.style, "eastAsia")
                if east in HEITI:
                    bad += 1
                    if len(samples) < 5:
                        samples.append(text[:16])
        self.report.metrics["body_heiti_runs"] = bad
        if bad < 3:
            return
        self.report.warn(
            "FONT-08",
            f"正文里有 {bad} 处文字是黑体（如“" + "”“".join(samples[:3]) + "”）。"
            "正文应为宋体；交叉引用（REF 域）更新时会继承题注的黑体，是最常见的原因",
            "把正文里的黑体改回宋体；交叉引用可在域代码后加 \\* MERGEFORMAT，"
            "让更新时保留域自身格式（见 references/cross-references.md）",
        )

    # -- 汇总 -------------------------------------------------------------
    def run(self) -> Report:
        self.check_pages()
        self.check_headers()
        self.check_front_matter()
        self.check_headings()
        self.check_abstract()
        self.check_body()
        self.check_latin_fonts()
        self.check_body_heiti()
        self.check_symbol_fonts()
        self.check_caption_layout()
        self.check_references()
        self.check_figures_tables()
        self.check_cross_references()
        return self.report


# ---------------------------------------------------------------------------
# 输出
# ---------------------------------------------------------------------------


def print_report(report: Report, chapter_mode: bool = False) -> None:
    print("=" * 72)
    print("湖南大学研究生学位论文格式自检报告")
    print("=" * 72)
    print(f"文件：{report.file}")
    print(f"学位类型：{DEGREES[report.degree]['label']}")
    if chapter_mode:
        print("模式：单章草稿（跳过封面/声明/摘要/目录/参考文献与总字数检查）")
    s = report.to_dict()["summary"]
    print(f"结果：ERROR {s['errors']} 项，WARN {s['warnings']} 项，INFO {s['infos']} 项")
    if report.metrics:
        print("\n关键指标：")
        labels = {
            "body_chinese_chars": "正文汉字数",
            "abstract_chars": "中文摘要汉字数",
            "keywords": "关键词个数",
            "references": "参考文献条数",
            "foreign_references": "外文文献条数（估）",
            "figures": "图题数",
            "tables": "表题数",
            "headings": "标题统计",
            "body_size_pt": "正文主字号(pt)",
            "body_font": "正文主字体",
            "body_line_spacing_pt": "正文固定行距(pt)",
            "in_text_citations": "正文引用标注数",
            "ref_fields": "REF 交叉引用域",
            "latin_cjk_runs": "西文误用中文字体（run）",
            "body_heiti_runs": "正文误用黑体（run）",
            "symbol_cjk_runs": "希腊字母/符号误用中文字体（run）",
        }
        for key, label in labels.items():
            if key in report.metrics:
                print(f"  - {label}：{report.metrics[key]}")

    groups = [
        ("ERROR", "错误（必须修）"),
        ("WARN", "警告（建议核对）"),
        ("INFO", "提示"),
    ]
    for level, title in groups:
        items = [i for i in report.issues if i.level == level]
        if not items:
            continue
        print(f"\n[{title}]")
        for idx, item in enumerate(items, 1):
            print(f"{idx:>2}. [{item.code}] {item.message}")
            if item.hint:
                print(f"     → {item.hint}")

    print("\n" + "-" * 72)
    if report.errors:
        print("结论：存在必须修复的 ERROR，暂不建议提交。")
    elif report.warnings:
        print("结论：无 ERROR；请逐条核对 WARN，确认是否有意保留。")
    else:
        print("结论：机器可查项目全部通过；请继续人工检查签名、印刷与内容质量。")
    print("-" * 72)


def main(argv: Optional[list[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="湖南大学学位论文格式自检")
    parser.add_argument("docx", help="待检查的 .docx 文件")
    parser.add_argument(
        "--degree",
        choices=sorted(DEGREES),
        default="master",
        help="学位类型：doctor / master / professional（默认 master）",
    )
    parser.add_argument("--json", metavar="FILE", help="把结果同时写入 JSON 文件")
    parser.add_argument(
        "--chapter",
        action="store_true",
        help="单章草稿模式：跳过封面/声明/摘要/目录/参考文献与总字数检查",
    )
    args = parser.parse_args(argv)

    path = Path(args.docx)
    if not path.exists():
        print(f"文件不存在：{path}", file=sys.stderr)
        return 2
    if path.suffix.lower() != ".docx":
        print("只支持 .docx；请先把 .doc 另存为 .docx", file=sys.stderr)
        return 2

    try:
        linter = Linter(path, args.degree, chapter_mode=args.chapter)
        report = linter.run()
    except Exception as exc:  # noqa: BLE001 - 给用户可读的错误
        print(f"读取失败：{exc}", file=sys.stderr)
        return 2

    print_report(report, chapter_mode=args.chapter)
    if args.json:
        Path(args.json).write_text(
            json.dumps(report.to_dict(), ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        print(f"\nJSON 报告已写入：{args.json}")
    return 1 if report.errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
